"""Parsing and validation for spreadsheet (xlsx / csv) product imports.

Pure functions, no DB access: the endpoint decides how to insert. Everything
here is synchronous and CPU-bound, so callers run it in a worker thread.
"""
import csv
import io
import re
from decimal import Decimal, InvalidOperation

MAX_FILE_BYTES = 10 * 1024 * 1024     # 10 MB upload cap
MAX_ROWS = 5000                       # data rows per file (keeps one request comfortably short)
MAX_SCANNED_ROWS = 50_000             # rows read incl. blanks, guards against sheets with a huge empty range

MAX_NAME = 300                        # products.name  String(300)
MAX_CATEGORY = 100                    # products.category String(100)
MAX_PRICE = Decimal("9999999999.99")  # products.price Numeric(12, 2)
MAX_STOCK = 2_000_000_000             # fits a 32-bit Integer

# canonical field -> accepted header spellings (compared after normalising)
_ALIASES = {
    "name": {"name", "productname", "product", "title", "producttitle"},
    "price": {"price", "unitprice", "sellingprice"},
    "description": {"description", "desc", "details"},
    "category": {"category", "categoryname"},
    "stock": {"stock", "quantity", "qty", "inventory"},
    "is_available": {"isavailable", "available", "active", "enabled"},
}
_IMAGE_RE = re.compile(r"^(image|images|img|photo|picture)(url|link)?\d*$")
_URL_RE = re.compile(r"https?://[^\s,;|]+", re.I)
_TRUE = {"true", "yes", "y", "1", "t", "available", "active"}
_FALSE = {"false", "no", "n", "0", "f", "unavailable", "inactive"}


class ImportFileError(Exception):
    """The file as a whole can't be used (wrong type, missing columns, too big...)."""


def _norm(header) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def _map_headers(header_row) -> tuple[dict, list]:
    """-> ({field: column index}, [image column indexes in sheet order])."""
    fields: dict[str, int] = {}
    images: list[int] = []
    for idx, h in enumerate(header_row):
        n = _norm(h)
        if not n:
            continue
        if _IMAGE_RE.match(n):
            images.append(idx)
            continue
        for field, names in _ALIASES.items():
            if n in names and field not in fields:
                fields[field] = idx
                break
    missing = [f for f in ("name", "price") if f not in fields]
    if missing:
        raise ImportFileError(
            "Missing required column(s): " + ", ".join(
                {"name": "Product Name", "price": "Price"}[m] for m in missing)
            + ". The first row of the file must be the header row.")
    return fields, images


def _read_csv(content: bytes):
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1252", errors="replace")   # Excel's "CSV" on Windows
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    for row in csv.reader(io.StringIO(text), dialect):
        yield row


def _read_xlsx(content: bytes):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Could not read this file. Make sure it is a valid .xlsx workbook.")
    try:
        ws = wb.worksheets[0]
        for row in ws.iter_rows(values_only=True):
            yield list(row)
    finally:
        wb.close()


def read_rows(filename: str, content: bytes):
    """Yield (sheet_row_number, cells) for every row of the file, header first."""
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        reader = _read_xlsx(content)
    elif name.endswith(".csv"):
        reader = _read_csv(content)
    elif name.endswith(".xls"):
        raise ImportFileError("Old .xls files aren't supported. Re-save the file as .xlsx or .csv.")
    else:
        raise ImportFileError("Unsupported file type. Upload a .xlsx or .csv file.")
    for n, cells in enumerate(reader, start=1):
        if n > MAX_SCANNED_ROWS:
            break
        yield n, cells


def _cell(cells, idx):
    if idx is None or idx >= len(cells):
        return None
    v = cells[idx]
    if isinstance(v, str):
        v = v.replace("\x00", "").strip()
        return v or None
    return v


def _clean_text(v, limit, label):
    s = str(v).replace("\x00", "").strip()
    if len(s) > limit:
        raise ValueError(f"{label} is longer than {limit} characters")
    return s


_THOUSANDS_RE = re.compile(r"^\d{1,3}(,\d{3})+(\.\d+)?$")


def _num_text(v) -> str:
    """'1,200.50' -> '1200.50'. A comma that isn't a thousands separator ('1,5')
    is left in place so the number is rejected instead of silently becoming 15."""
    s = str(v).strip().replace("$", "").replace(" ", "")
    return s.replace(",", "") if _THOUSANDS_RE.match(s) else s


def _parse_price(v) -> Decimal:
    if v is None:
        raise ValueError("price is required")
    s = _num_text(v)
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"price '{v}' is not a number")
    if not d.is_finite() or d < 0:
        raise ValueError("price must be 0 or more")
    d = d.quantize(Decimal("0.01"))
    if d > MAX_PRICE:
        raise ValueError("price is too large")
    return d


def _parse_stock(v) -> int:
    if v is None:
        return 0
    s = _num_text(v)
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"stock '{v}' is not a number")
    if not d.is_finite() or d < 0 or d != d.to_integral_value():
        raise ValueError("stock must be a whole number, 0 or more")
    if d > MAX_STOCK:
        raise ValueError("stock is too large")
    return int(d)


def _parse_bool(v) -> bool:
    if v is None:
        return True                      # blank = available
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    if s in _TRUE:
        return True
    if s in _FALSE:
        return False
    raise ValueError(f"is_available '{v}' must be True/False")


def _first_image(cells, image_cols):
    """The system shows one image per product: use the first valid link found
    across the image columns (a cell may itself hold several links)."""
    for idx in image_cols:
        v = _cell(cells, idx)
        if v:
            m = _URL_RE.search(str(v))
            if m:
                return m.group(0)
    return None


def parse_products(filename: str, content: bytes):
    """-> (valid, errors, total_rows).

    valid:  [{row, name, price, category, stock, image, description, is_available}]
    errors: [{row, name, reason}]   row = the row number as seen in Excel
    """
    if len(content) > MAX_FILE_BYTES:
        raise ImportFileError(f"File is too large (max {MAX_FILE_BYTES // (1024 * 1024)} MB).")

    rows = read_rows(filename, content)
    header = None
    fields, images = {}, []
    valid, errors = [], []
    total = 0
    for n, cells in rows:
        if header is None:
            if not any(c not in (None, "") for c in cells):
                continue                  # leading blank rows
            header = cells
            fields, images = _map_headers(header)
            continue
        if not any(c not in (None, "") for c in cells):
            continue                      # skip blank rows
        total += 1
        if total > MAX_ROWS:
            raise ImportFileError(
                f"Too many rows: the limit is {MAX_ROWS} products per file. Split the file and upload it in parts.")
        raw_name = _cell(cells, fields["name"])
        try:
            if not raw_name:
                raise ValueError("product name is required")
            name = _clean_text(raw_name, MAX_NAME, "product name")
            price = _parse_price(_cell(cells, fields["price"]))
            stock = _parse_stock(_cell(cells, fields.get("stock")))
            category = _clean_text(_cell(cells, fields.get("category")) or "General", MAX_CATEGORY, "category")
            is_available = _parse_bool(_cell(cells, fields.get("is_available")))
            desc = _cell(cells, fields.get("description"))
            valid.append({
                "row": n, "name": name, "price": price, "category": category, "stock": stock,
                "image": _first_image(cells, images),
                "description": str(desc) if desc is not None else None,
                "is_available": is_available,
            })
        except ValueError as e:
            errors.append({"row": n, "name": str(raw_name)[:80] if raw_name else None, "reason": str(e)})

    if header is None:
        raise ImportFileError("The file is empty.")
    if total == 0:
        raise ImportFileError("No product rows found under the header row.")
    return valid, errors, total
